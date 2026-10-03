from flask import Flask, request, jsonify, send_file, render_template
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from pathlib import Path
from threading import Lock
from datetime import datetime
import os

app = Flask(__name__)
BASE = Path(__file__).resolve().parent
# Persistent data location. Locally this remains ./data; on Render we use /var/data
# so a Render Persistent Disk can preserve the Excel workbook across restarts/deploys.
DATA_DIR = Path(os.environ.get("DATA_DIR", str(BASE / "data")))
EXCEL_FILE = DATA_DIR / "Service_Complaint_Tracker.xlsx"
TEMPLATE_FILE = BASE / "data" / "Service_Complaint_Tracker.xlsx"

DATA_DIR.mkdir(parents=True, exist_ok=True)
if not EXCEL_FILE.exists() and TEMPLATE_FILE.exists() and EXCEL_FILE != TEMPLATE_FILE:
    import shutil
    shutil.copy2(TEMPLATE_FILE, EXCEL_FILE)
SHEET = "Service Tracker"
LOCK = Lock()

HEADERS = [
    "Sl. No.", "Inward Date", "Outward Date", "Model No.",
    "Customer Details", "Complaints", "Status", "Remarks",
    "Item Replaced", "Final Remark", "Service Charge"
]

def find_next_row(ws):
    # Find the first completely empty service-data row.
    for r in range(2, ws.max_row + 2):
        values = [ws.cell(r, c).value for c in range(2, 12)]
        if all(v in (None, "") for v in values):
            return r
    return ws.max_row + 1

def ensure_setup(wb, ws):
    # Keep the user's existing workbook formatting and sheets.
    # Extend status dropdown/conditional formatting for future rows.
    dv = DataValidation(type="list", formula1='"Complete,Pending"', allow_blank=True)
    dv.error = "Please select Complete or Pending."
    dv.errorTitle = "Invalid Status"
    ws.add_data_validation(dv)
    dv.add("G2:G10000")

    green_fill = PatternFill("solid", fgColor="C6EFCE")
    green_font = Font(color="006100", bold=True)
    red_fill = PatternFill("solid", fgColor="FFC7CE")
    red_font = Font(color="9C0006", bold=True)
    ws.conditional_formatting.add(
        "G2:G10000",
        FormulaRule(formula=['G2="Complete"'], fill=green_fill, font=green_font)
    )
    ws.conditional_formatting.add(
        "G2:G10000",
        FormulaRule(formula=['G2="Pending"'], fill=red_fill, font=red_font)
    )

def update_summary(wb, last_row):
    if "Summary" not in wb.sheetnames:
        return
    s = wb["Summary"]
    s["B3"] = f'=COUNTA(\'{SHEET}\'!D2:D{last_row})'
    s["B4"] = f'=COUNTIF(\'{SHEET}\'!G2:G{last_row},"Complete")'
    s["B5"] = f'=COUNTIF(\'{SHEET}\'!G2:G{last_row},"Pending")'
    s["B6"] = f'=SUM(\'{SHEET}\'!K2:K{last_row})'

@app.get("/")
def home():
    return render_template("index.html")

@app.get("/api/count")
def count():
    with LOCK:
        wb = load_workbook(EXCEL_FILE, read_only=True, data_only=False)
        ws = wb[SHEET]
        n = 0
        for r in range(2, ws.max_row + 1):
            if any(ws.cell(r,c).value not in (None, "") for c in range(2,12)):
                n += 1
        wb.close()
    return jsonify(count=n)

@app.post("/api/entries")
def add_entry():
    data = request.get_json(silent=True) or {}
    with LOCK:
        wb = load_workbook(EXCEL_FILE)
        ws = wb[SHEET]

        # Verify the expected header layout.
        actual = [ws.cell(1,c).value for c in range(1,12)]
        if actual != HEADERS:
            wb.close()
            return jsonify(error="The Excel column layout does not match the service tracker."), 400

        row = find_next_row(ws)

        # Sl. No. is generated automatically.
        previous_numbers = []
        for r in range(2, row):
            v = ws.cell(r,1).value
            if isinstance(v, (int, float)):
                previous_numbers.append(int(v))
        slno = max(previous_numbers, default=0) + 1

        values = [
            slno,
            data.get("inwardDate", ""),
            data.get("outwardDate", ""),
            data.get("modelNo", "").strip(),
            data.get("customerDetails", "").strip(),
            data.get("complaints", "").strip(),
            "Complete" if data.get("status") == "Complete" else "Pending",
            data.get("remarks", "").strip(),
            data.get("itemReplaced", "").strip(),
            data.get("finalRemark", "").strip(),
            data.get("serviceCharge", "")
        ]

        for c, value in enumerate(values, 1):
            ws.cell(row, c).value = value
            ws.cell(row, c).alignment = Alignment(vertical="top", wrap_text=True)

        # Keep date/charge formatting.
        ws.cell(row, 2).number_format = "dd-mm-yyyy"
        ws.cell(row, 3).number_format = "dd-mm-yyyy"
        ws.cell(row, 11).number_format = '₹#,##0.00'

        # Extend the existing filter range.
        ws.auto_filter.ref = f"A1:K{max(row, 501)}"
        ws.freeze_panes = "A2"

        ensure_setup(wb, ws)
        update_summary(wb, row)

        # Ensure workbook calculations refresh when opened in Excel.
        try:
            wb.calculation.fullCalcOnLoad = True
            wb.calculation.forceFullCalc = True
        except Exception:
            pass

        wb.save(EXCEL_FILE)
        wb.close()

    return jsonify(ok=True, slno=slno, row=row)

@app.post("/api/clear")
def clear_excel():
    # Delete all service entries while keeping the workbook, headers, formatting,
    # dropdowns, conditional formatting, and summary sheet structure.
    with LOCK:
        wb = load_workbook(EXCEL_FILE)
        if SHEET not in wb.sheetnames:
            wb.close()
            return jsonify(error="Service Tracker sheet not found."), 400
        ws = wb[SHEET]

        # Clear service data rows, but keep row 1 (headers) and formatting.
        for r in range(2, ws.max_row + 1):
            for c in range(1, 12):
                ws.cell(r, c).value = None

        # Keep the existing summary sheet and reset its formulas to zero.
        if "Summary" in wb.sheetnames:
            summary = wb["Summary"]
            summary["B3"] = 0
            summary["B4"] = 0
            summary["B5"] = 0
            summary["B6"] = 0

        try:
            wb.calculation.fullCalcOnLoad = True
            wb.calculation.forceFullCalc = True
        except Exception:
            pass

        wb.save(EXCEL_FILE)
        wb.close()

    return jsonify(ok=True)

@app.get("/api/download")
def download():
    return send_file(EXCEL_FILE, as_attachment=True, download_name="Service_Complaint_Tracker.xlsx")

if __name__ == "__main__":
    port = int(os.environ.get("PORT", "3000"))
    print(f"Service Entry App running at http://127.0.0.1:{port}")
    app.run(host="0.0.0.0", port=port, debug=False)
